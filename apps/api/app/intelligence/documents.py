"""Document extraction with page and sheet provenance.

PyMuPDF is AGPL and is not imported. PDF text comes from pypdf (BSD).
DOCX uses python-docx. XLSX uses openpyxl.
"""

from __future__ import annotations

import csv
import io
from html.parser import HTMLParser
from pathlib import Path


class _Text(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        if data and data.strip():
            self.parts.append(data.strip())


def extract(path: str | None = None, data: bytes | None = None, name: str = "") -> dict:
    raw_name = name or (path or "")
    suffix = Path(raw_name).suffix.lower()
    blob = data if data is not None else (Path(path).read_bytes() if path else b"")
    if suffix == ".pdf":
        return _pdf(blob, raw_name)
    if suffix == ".docx":
        return _docx(blob, raw_name)
    if suffix == ".xlsx":
        return _xlsx(blob, raw_name)
    if suffix == ".csv":
        text = blob.decode("utf-8", errors="replace")
        rows = list(csv.reader(io.StringIO(text)))
        rendered = "\n".join(" | ".join(row) for row in rows[:40])
        return _ok(rendered, raw_name, "csv", tables=[{"sheet": "csv", "rows": rows[:40]}])
    if suffix in {".html", ".htm"}:
        parser = _Text()
        parser.feed(blob.decode("utf-8", errors="replace"))
        return _ok("\n".join(parser.parts), raw_name, "html")
    if suffix in {".txt", ".md", ""}:
        text = blob.decode("utf-8", errors="replace")
        return _ok(text, raw_name, "text")
    return {"ok": False, "kind": suffix or "unknown", "text": "", "error": "unsupported_document", "source": raw_name}


def _pdf(blob: bytes, source: str) -> dict:
    try:
        from pypdf import PdfReader
    except Exception as exc:
        return {"ok": False, "kind": "pdf", "text": "", "error": f"pypdf_import_failed: {type(exc).__name__}", "source": source}
    try:
        reader = PdfReader(io.BytesIO(blob))
    except Exception as exc:
        return {"ok": False, "kind": "pdf", "text": "", "error": f"pdf_read_failed: {type(exc).__name__}", "source": source}
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        pages.append({"page": index, "text": page.extract_text() or ""})
    text = "\n".join(f"[page {item['page']}] {item['text']}" for item in pages)
    return _ok(text, source, "pdf", pages=pages, page=1 if pages else None)


def _docx(blob: bytes, source: str) -> dict:
    try:
        from docx import Document
    except Exception as exc:
        return {"ok": False, "kind": "docx", "text": "", "error": f"docx_import_failed: {type(exc).__name__}", "source": source}
    document = Document(io.BytesIO(blob))
    paragraphs = [{"paragraph": index, "text": para.text} for index, para in enumerate(document.paragraphs, start=1) if para.text.strip()]
    text = "\n".join(f"[paragraph {item['paragraph']}] {item['text']}" for item in paragraphs)
    return _ok(text, source, "docx", paragraphs=paragraphs)


def _xlsx(blob: bytes, source: str) -> dict:
    try:
        from openpyxl import load_workbook
    except Exception as exc:
        return {"ok": False, "kind": "xlsx", "text": "", "error": f"openpyxl_import_failed: {type(exc).__name__}", "source": source}
    book = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    tables = []
    lines = []
    for sheet in book.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(["" if cell is None else str(cell) for cell in row])
        tables.append({"sheet": sheet.title, "rows": rows[:40]})
        lines.append(f"[sheet {sheet.title}]")
        lines.extend(" | ".join(row) for row in rows[:40])
    book.close()
    return _ok("\n".join(lines), source, "xlsx", tables=tables, sheet=tables[0]["sheet"] if tables else "")


def _ok(text: str, source: str, kind: str, tables: list | None = None, pages: list | None = None, paragraphs: list | None = None, page=None, sheet: str = "") -> dict:
    return {
        "ok": True,
        "kind": kind,
        "text": text[:8000],
        "tables": tables or [],
        "pages": pages or [],
        "paragraphs": paragraphs or [],
        "source": source,
        "page": page,
        "sheet": sheet,
    }


def status() -> dict:
    return {
        "backend": "pypdf+python-docx+openpyxl",
        "formats": ["txt", "md", "csv", "html", "pdf", "docx", "xlsx"],
        "pdf": "ACTIVE",
        "pymupdf": "REJECTED_AGPL",
        "pypdf": "INTEGRATED",
        "python_docx": "INTEGRATED",
        "openpyxl": "INTEGRATED",
    }
